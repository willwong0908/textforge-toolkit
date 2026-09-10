import inspect
import json
import re
from collections import Counter
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook as real_load_workbook

from term_extractor_app.constants import TERM_LIBRARY_SHEET
from term_extractor_app.service_layer import ExtractionTaskFacade
from term_extractor_app.web_app import APP_JS, build_index_html, create_app
import webui_launcher


def test_blocking_desktop_and_file_routes_run_in_fastapi_threadpool() -> None:
    app = create_app()
    blocking_paths = {
        "/api/dialog/select-folder",
        "/api/dialog/select-review-file",
        "/api/dialog/select-review-files",
        "/api/dialog/select-preprocess-files",
        "/api/preprocess/file-mapping-scan",
        "/api/cross-excel/search",
        "/api/cross-excel/merge",
        "/api/results/summary",
    }
    endpoints = {route.path: route.endpoint for route in app.routes if route.path in blocking_paths}

    assert endpoints.keys() == blocking_paths
    assert all(not inspect.iscoroutinefunction(endpoint) for endpoint in endpoints.values())


def test_result_summary_reuses_workbook_until_output_changes(tmp_path: Path, monkeypatch) -> None:
    output_path = tmp_path / "result.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = TERM_LIBRARY_SHEET
    sheet.append(["term"])
    sheet.append(["alpha"])
    workbook.save(output_path)

    calls = 0

    def counting_load_workbook(*args, **kwargs):
        nonlocal calls
        calls += 1
        return real_load_workbook(*args, **kwargs)

    monkeypatch.setattr("term_extractor_app.service_layer.load_workbook", counting_load_workbook)
    facade = ExtractionTaskFacade()

    first = facade.result_summary(str(output_path))
    second = facade.result_summary(str(output_path))

    assert first.term_library_count == 1
    assert second.term_library_count == 1
    assert calls == 1

    changed = real_load_workbook(output_path)
    changed[TERM_LIBRARY_SHEET].append(["beta"])
    changed.save(output_path)
    changed.close()

    third = facade.result_summary(str(output_path))
    assert third.term_library_count == 2
    assert calls == 2


def test_frontend_polling_has_timeouts_and_overlap_guards() -> None:
    assert "timeoutMs = 60000" in APP_JS
    assert "statusRefreshInFlight" in APP_JS
    assert "aiReviewTaskPollInFlight" in APP_JS
    assert "setInterval(refreshStatusSafely, 1500)" in APP_JS
    assert "data.output_file !== lastResultSummaryPath" in APP_JS


def test_static_ui_has_unique_ids_and_no_unbound_id_buttons() -> None:
    html = build_index_html()
    ids = re.findall(r'\bid="([^"]+)"', html)
    duplicates = [element_id for element_id, count in Counter(ids).items() if count > 1]
    button_ids = re.findall(r'<button\b[^>]*\bid="([^"]+)"[^>]*>', html, re.IGNORECASE)
    # This button is a native method="dialog" submit control and deliberately needs no JS binding.
    intentionally_native = {"closeReviewAttachmentMappingButton"}
    unbound = [button_id for button_id in button_ids if button_id not in APP_JS and button_id not in intentionally_native]

    assert duplicates == []
    assert unbound == []


def test_main_page_still_serves_after_resilience_changes() -> None:
    client = TestClient(create_app())
    response = client.get("/")
    assert response.status_code == 200
    assert "译禾工具合集" in response.text
    assert client.get("/favicon.ico").status_code == 204


class _LauncherResponse:
    status = 200

    def __init__(self, payload: object) -> None:
        self.payload = payload

    def __enter__(self):
        return self

    def __exit__(self, *_args) -> None:
        return None

    def read(self) -> bytes:
        return json.dumps(self.payload).encode("utf-8")


def test_launcher_only_accepts_the_expected_http_service(monkeypatch) -> None:
    monkeypatch.setattr(
        webui_launcher.urllib.request,
        "urlopen",
        lambda *_args, **_kwargs: _LauncherResponse({"is_running": False, "stage": "idle"}),
    )
    assert webui_launcher._is_service_ready() is True

    monkeypatch.setattr(
        webui_launcher.urllib.request,
        "urlopen",
        lambda *_args, **_kwargs: _LauncherResponse({"unrelated": "service"}),
    )
    assert webui_launcher._is_service_ready() is False


def test_launcher_rejects_an_unresponsive_port(monkeypatch) -> None:
    def fail(*_args, **_kwargs):
        raise TimeoutError("not an HTTP service")

    monkeypatch.setattr(webui_launcher.urllib.request, "urlopen", fail)
    assert webui_launcher._is_service_ready() is False
