from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET

from openpyxl.utils import get_column_letter

from ..excel_streaming import open_streaming_workbook
from .types import ReaderBlock, ReaderDocument
from .xliff_reader import read_xliff_items, read_xliff_language_metadata


class ReaderError(ValueError):
    pass


class ReaderDependencyError(ReaderError):
    def __init__(self, message: str, *, package_name: str, source_url: str) -> None:
        super().__init__(message)
        self.package_name = package_name
        self.source_url = source_url


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _role_hint(filename: str, labels: Iterable[str]) -> str:
    combined = " ".join([filename, *labels]).lower()
    reference_markers = (
        "glossary", "term", "style", "guide", "reference", "context", "说明", "术语", "规范", "参考",
    )
    return "reference" if any(marker in combined for marker in reference_markers) else "content"


class BaseReader(ABC):
    name = "base"
    extensions: tuple[str, ...] = ()

    @abstractmethod
    def read(self, path: Path, original_filename: str) -> ReaderDocument:
        raise NotImplementedError


class ExcelReader(BaseReader):
    name = "excel"
    extensions = (".xlsx", ".xlsm")

    def read(self, path: Path, original_filename: str) -> ReaderDocument:
        blocks: list[ReaderBlock] = []
        sheets: list[dict[str, Any]] = []
        with open_streaming_workbook(path) as workbook:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header_row = next(rows, None) or ()
                headers = [_text(value) for value in header_row]
                used_columns: set[int] = {index for index, value in enumerate(headers) if value}
                row_count = 1
                for row_number, row in enumerate(rows, start=2):
                    row_count = row_number
                    for column_index, value in enumerate(row):
                        value_text = _text(value)
                        if not value_text:
                            continue
                        used_columns.add(column_index)
                        letter = get_column_letter(column_index + 1)
                        header = headers[column_index] if column_index < len(headers) else ""
                        blocks.append(
                            ReaderBlock(
                                pointer=f"sheet:{sheet.title}/cell:{letter}{row_number}",
                                text=value_text,
                                label=header or letter,
                                metadata={
                                    "sheet": sheet.title,
                                    "row": row_number,
                                    "column_index": column_index,
                                    "column": letter,
                                    "header": header,
                                },
                            )
                        )
                sheets.append(
                    {
                        "name": sheet.title,
                        "row_count": row_count,
                        "columns": [
                            {
                                "index": index,
                                "letter": get_column_letter(index + 1),
                                "header": headers[index] if index < len(headers) else "",
                            }
                            for index in sorted(used_columns)
                        ],
                    }
                )
        if not blocks:
            raise ReaderError("Excel 中没有可读取的正文单元格")
        labels = [column["header"] for sheet in sheets for column in sheet["columns"]]
        return ReaderDocument(
            reader_name=self.name,
            file_type="excel",
            filename=original_filename,
            file_hash=file_sha256(path),
            role_hint=_role_hint(original_filename, labels),
            structure={"sheets": sheets},
            blocks=blocks,
        )


class XliffReader(BaseReader):
    name = "xliff"
    extensions = (".xlf", ".xliff")

    def read(self, path: Path, original_filename: str) -> ReaderDocument:
        items = read_xliff_items(path, original_filename)
        metadata = read_xliff_language_metadata(path)
        blocks = [
            ReaderBlock(
                pointer=f"segment:{item.get('segment_id') or item.get('item_order')}",
                text=str(item.get("target_text") or ""),
                label="target",
                language_hint=str(metadata.get("target_language") or ""),
                metadata={
                    "source_text": str(item.get("source_text") or ""),
                    "source_language": str(metadata.get("source_language") or ""),
                    "target_language": str(metadata.get("target_language") or ""),
                    "segment_id": item.get("segment_id"),
                    "row_number": item.get("row_number"),
                },
            )
            for item in items
            if str(item.get("target_text") or "").strip()
        ]
        if not blocks:
            raise ReaderError("XLIFF 中没有可审校译文")
        return ReaderDocument(
            reader_name=self.name,
            file_type="xliff",
            filename=original_filename,
            file_hash=file_sha256(path),
            structure={"segment_count": len(items), **metadata},
            blocks=blocks,
        )


class DelimitedReader(BaseReader):
    name = "delimited"
    extensions = (".csv", ".tsv")

    def read(self, path: Path, original_filename: str) -> ReaderDocument:
        raw = path.read_bytes()
        text, encoding = decode_text(raw)
        delimiter = "\t" if path.suffix.lower() == ".tsv" else _sniff_delimiter(text)
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        if not rows:
            raise ReaderError("表格文件为空")
        headers = [_text(value) for value in rows[0]]
        blocks: list[ReaderBlock] = []
        for row_number, row in enumerate(rows[1:], start=2):
            for column_index, value in enumerate(row):
                value_text = _text(value)
                if not value_text:
                    continue
                label = headers[column_index] if column_index < len(headers) else str(column_index + 1)
                blocks.append(
                    ReaderBlock(
                        pointer=f"row:{row_number}/column:{column_index + 1}",
                        text=value_text,
                        label=label,
                        metadata={"row": row_number, "column_index": column_index, "header": label},
                    )
                )
        if not blocks:
            raise ReaderError("表格文件没有可读取的正文")
        return ReaderDocument(
            reader_name=self.name,
            file_type=path.suffix.lower().lstrip("."),
            filename=original_filename,
            file_hash=file_sha256(path),
            role_hint=_role_hint(original_filename, headers),
            structure={"encoding": encoding, "delimiter": delimiter, "headers": headers, "row_count": len(rows)},
            blocks=blocks,
        )


class TextReader(BaseReader):
    name = "text"
    extensions = (".txt", ".md", ".log")

    def read(self, path: Path, original_filename: str) -> ReaderDocument:
        text, encoding = decode_text(path.read_bytes())
        chunks = [chunk.strip() for chunk in re.split(r"\n\s*\n|\n", text) if chunk.strip()]
        if not chunks:
            raise ReaderError("文本文件为空")
        return ReaderDocument(
            reader_name=self.name,
            file_type="text",
            filename=original_filename,
            file_hash=file_sha256(path),
            role_hint=_role_hint(original_filename, []),
            structure={"encoding": encoding, "paragraph_count": len(chunks)},
            blocks=[ReaderBlock(pointer=f"paragraph:{index}", text=value) for index, value in enumerate(chunks, 1)],
        )


class JsonReader(BaseReader):
    name = "json"
    extensions = (".json",)

    def read(self, path: Path, original_filename: str) -> ReaderDocument:
        text, encoding = decode_text(path.read_bytes())
        try:
            data = json.loads(text)
        except json.JSONDecodeError as exc:
            raise ReaderError(f"JSON 解析失败：{exc}") from exc
        blocks = [ReaderBlock(pointer=pointer, text=value, label=pointer.rsplit("/", 1)[-1]) for pointer, value in _json_leaves(data)]
        if not blocks:
            raise ReaderError("JSON 中没有可读取的文本值")
        return ReaderDocument(
            reader_name=self.name,
            file_type="json",
            filename=original_filename,
            file_hash=file_sha256(path),
            role_hint=_role_hint(original_filename, [block.label for block in blocks[:50]]),
            structure={"encoding": encoding, "root_type": type(data).__name__, "text_value_count": len(blocks)},
            blocks=blocks,
        )


class XmlReader(BaseReader):
    name = "xml"
    extensions = (".xml",)

    def read(self, path: Path, original_filename: str) -> ReaderDocument:
        try:
            root = ET.parse(path).getroot()
        except ET.ParseError as exc:
            raise ReaderError(f"XML 解析失败：{exc}") from exc
        blocks: list[ReaderBlock] = []
        counters: dict[str, int] = {}
        for element in root.iter():
            value = _text(element.text)
            if not value:
                continue
            tag = element.tag.rsplit("}", 1)[-1]
            counters[tag] = counters.get(tag, 0) + 1
            blocks.append(ReaderBlock(pointer=f"element:{tag}[{counters[tag]}]", text=value, label=tag))
        if not blocks:
            raise ReaderError("XML 中没有可读取的文本节点")
        return ReaderDocument(
            reader_name=self.name,
            file_type="xml",
            filename=original_filename,
            file_hash=file_sha256(path),
            role_hint=_role_hint(original_filename, counters.keys()),
            structure={"root_tag": root.tag.rsplit("}", 1)[-1], "tag_counts": counters},
            blocks=blocks,
        )


class OfficeXmlReader(BaseReader):
    def _paragraphs(self, path: Path, member_pattern: re.Pattern[str]) -> list[tuple[str, str]]:
        paragraphs: list[tuple[str, str]] = []
        with zipfile.ZipFile(path) as archive:
            members = sorted(name for name in archive.namelist() if member_pattern.fullmatch(name))
            for member in members:
                root = ET.fromstring(archive.read(member))
                paragraph_index = 0
                for paragraph in root.iter():
                    if paragraph.tag.rsplit("}", 1)[-1] != "p":
                        continue
                    value = "".join(node.text or "" for node in paragraph.iter() if node.tag.rsplit("}", 1)[-1] == "t").strip()
                    if value:
                        paragraph_index += 1
                        paragraphs.append((f"{member}/paragraph:{paragraph_index}", value))
        return paragraphs


class DocxReader(OfficeXmlReader):
    name = "docx"
    extensions = (".docx",)

    def read(self, path: Path, original_filename: str) -> ReaderDocument:
        paragraphs = self._paragraphs(path, re.compile(r"word/(document|header\d+|footer\d+)\.xml"))
        if not paragraphs:
            raise ReaderError("DOCX 中没有可读取的文本")
        return ReaderDocument(
            reader_name=self.name,
            file_type="docx",
            filename=original_filename,
            file_hash=file_sha256(path),
            role_hint=_role_hint(original_filename, []),
            structure={"paragraph_count": len(paragraphs)},
            blocks=[ReaderBlock(pointer=pointer, text=value) for pointer, value in paragraphs],
        )


class PptxReader(OfficeXmlReader):
    name = "pptx"
    extensions = (".pptx",)

    def read(self, path: Path, original_filename: str) -> ReaderDocument:
        paragraphs = self._paragraphs(path, re.compile(r"ppt/slides/slide\d+\.xml"))
        if not paragraphs:
            raise ReaderError("PPTX 中没有可读取的文本")
        return ReaderDocument(
            reader_name=self.name,
            file_type="pptx",
            filename=original_filename,
            file_hash=file_sha256(path),
            role_hint=_role_hint(original_filename, []),
            structure={"paragraph_count": len(paragraphs)},
            blocks=[ReaderBlock(pointer=pointer, text=value) for pointer, value in paragraphs],
        )


class PdfReader(BaseReader):
    name = "pdf"
    extensions = (".pdf",)

    def read(self, path: Path, original_filename: str) -> ReaderDocument:
        try:
            from pypdf import PdfReader as PyPdfReader
        except ImportError as exc:
            raise ReaderDependencyError(
                "读取 PDF 需要安装 pypdf",
                package_name="pypdf",
                source_url="https://pypi.org/project/pypdf/",
            ) from exc
        pdf = PyPdfReader(str(path))
        blocks: list[ReaderBlock] = []
        for page_index, page in enumerate(pdf.pages, 1):
            page_text = str(page.extract_text() or "")
            for paragraph_index, value in enumerate((part.strip() for part in page_text.splitlines() if part.strip()), 1):
                blocks.append(ReaderBlock(pointer=f"page:{page_index}/paragraph:{paragraph_index}", text=value))
        if not blocks:
            raise ReaderError("PDF 没有可提取文本，可能是扫描件；当前版本不包含 OCR")
        return ReaderDocument(
            reader_name=self.name,
            file_type="pdf",
            filename=original_filename,
            file_hash=file_sha256(path),
            role_hint=_role_hint(original_filename, []),
            structure={"page_count": len(pdf.pages)},
            blocks=blocks,
        )


class ReaderRegistry:
    def __init__(self) -> None:
        self._readers: dict[str, BaseReader] = {}

    def register(self, reader: BaseReader) -> None:
        for extension in reader.extensions:
            self._readers[extension.lower()] = reader

    def reader_for(self, path: Path) -> BaseReader:
        reader = self._readers.get(path.suffix.lower())
        if reader is None:
            raise ReaderError(f"暂不支持 {path.suffix or '无扩展名'} 文件，需要安装读取器")
        return reader

    def read(self, path: Path, original_filename: str | None = None) -> ReaderDocument:
        return self.reader_for(path).read(path, original_filename or path.name)

    @property
    def supported_extensions(self) -> tuple[str, ...]:
        return tuple(sorted(self._readers))


def build_default_registry() -> ReaderRegistry:
    registry = ReaderRegistry()
    for reader in (
        ExcelReader(), XliffReader(), DelimitedReader(), TextReader(), JsonReader(), XmlReader(),
        DocxReader(), PptxReader(), PdfReader(),
    ):
        registry.register(reader)
    return registry


def decode_text(raw: bytes) -> tuple[str, str]:
    if not raw:
        return "", "utf-8"
    try:
        return raw.decode("utf-8-sig"), "utf-8"
    except UnicodeDecodeError:
        try:
            from charset_normalizer import from_bytes

            best = from_bytes(raw).best()
            if best is not None:
                return str(best), str(best.encoding or "unknown")
        except Exception:
            pass
    for encoding in ("gb18030", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise ReaderError("无法识别文本编码")


def _sniff_delimiter(text: str) -> str:
    try:
        return csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _json_leaves(value: Any, pointer: str = "$") -> Iterable[tuple[str, str]]:
    if isinstance(value, dict):
        for key, child in value.items():
            yield from _json_leaves(child, f"{pointer}/{key}")
    elif isinstance(value, list):
        for index, child in enumerate(value):
            yield from _json_leaves(child, f"{pointer}/{index}")
    elif isinstance(value, str) and value.strip():
        yield pointer, value.strip()
