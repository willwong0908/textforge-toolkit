from __future__ import annotations

import hashlib
import re
import unicodedata
import uuid
from collections import OrderedDict
from io import BytesIO
from pathlib import Path
from threading import RLock
from typing import Any

from openpyxl import load_workbook

from .config import TERM_BASES_DIR, ensure_directories
from .database import dumps_json, get_connection, init_db, loads_json, utc_now


MAX_TERM_BASE_BYTES = 50 * 1024 * 1024
MAX_MATCHED_ENTRIES_PER_ITEM = 50
SUPPORTED_TERM_BASE_EXTENSIONS = {".xlsx", ".xlsm"}
METADATA_HEADERS = {
    "entry_id": "entry_id",
    "entry_subject": "entry_subject",
    "entry_note": "entry_note",
}

# UI language names, locale codes, and terminology-table headers all resolve to
# an exact language column. Generic languages prefer the generic column and only
# fall back to a regional column if the generic column is absent.
LANGUAGE_ALIASES: dict[str, tuple[str, ...]] = {
    "chineseprc": ("Chinese_PRC",),
    "chinesesimplified": ("Chinese_PRC",),
    "simplifiedchinese": ("Chinese_PRC",),
    "简体中文": ("Chinese_PRC",),
    "中文": ("Chinese_PRC", "Chinese_Taiwan"),
    "zh": ("Chinese_PRC", "Chinese_Taiwan"),
    "zhcn": ("Chinese_PRC",),
    "zhhans": ("Chinese_PRC",),
    "chinesetaiwan": ("Chinese_Taiwan",),
    "chinesetraditional": ("Chinese_Taiwan",),
    "traditionalchinese": ("Chinese_Taiwan",),
    "繁体中文": ("Chinese_Taiwan",),
    "zhtw": ("Chinese_Taiwan",),
    "zhhant": ("Chinese_Taiwan",),
    "englishunitedkingdom": ("English_United_Kingdom", "English"),
    "engb": ("English_United_Kingdom", "English"),
    "english": ("English", "English_United_Kingdom"),
    "英语": ("English", "English_United_Kingdom"),
    "en": ("English", "English_United_Kingdom"),
    "enus": ("English", "English_United_Kingdom"),
    "japanese": ("Japanese",), "日语": ("Japanese",), "ja": ("Japanese",),
    "korean": ("Korean",), "韩语": ("Korean",), "ko": ("Korean",),
    "french": ("French",), "法语": ("French",), "fr": ("French",),
    "german": ("German",), "德语": ("German",), "de": ("German",),
    "italian": ("Italian",), "意大利语": ("Italian",), "it": ("Italian",),
    "spanish": ("Spanish",), "西班牙语": ("Spanish",), "es": ("Spanish",),
    "portuguese": ("Portuguese",), "葡萄牙语": ("Portuguese",), "pt": ("Portuguese",),
    "russian": ("Russian",), "俄语": ("Russian",), "ru": ("Russian",),
    "arabic": ("Arabic",), "阿拉伯语": ("Arabic",), "ar": ("Arabic",),
    "thai": ("Thai",), "泰语": ("Thai",), "th": ("Thai",),
    "vietnamese": ("Vietnamese",), "越南语": ("Vietnamese",), "vi": ("Vietnamese",),
    "indonesian": ("Indonesian",), "印尼语": ("Indonesian",), "id": ("Indonesian",),
    "turkish": ("Turkish",), "土耳其语": ("Turkish",), "tr": ("Turkish",),
    "polish": ("Polish",), "波兰语": ("Polish",), "pl": ("Polish",),
    "dutch": ("Dutch",), "荷兰语": ("Dutch",), "nl": ("Dutch",),
    "swedish": ("Swedish",), "瑞典语": ("Swedish",), "sv": ("Swedish",),
    "norwegian": ("Norwegian",), "挪威语": ("Norwegian",), "no": ("Norwegian",),
    "danish": ("Danish",), "丹麦语": ("Danish",), "da": ("Danish",),
    "finnish": ("Finnish",), "芬兰语": ("Finnish",), "fi": ("Finnish",),
    "czech": ("Czech",), "捷克语": ("Czech",), "cs": ("Czech",),
    "hungarian": ("Hungarian",), "匈牙利语": ("Hungarian",), "hu": ("Hungarian",),
    "romanian": ("Romanian",), "罗马尼亚语": ("Romanian",), "ro": ("Romanian",),
    "greek": ("Greek",), "希腊语": ("Greek",), "el": ("Greek",),
    "hebrew": ("Hebrew",), "希伯来语": ("Hebrew",), "he": ("Hebrew",),
    "ukrainian": ("Ukrainian",), "乌克兰语": ("Ukrainian",), "uk": ("Ukrainian",),
}

_matcher_cache: OrderedDict[tuple[str, str, str, str], dict[str, Any]] = OrderedDict()
_matcher_lock = RLock()


class TermBaseError(ValueError):
    pass


def list_term_bases() -> list[dict[str, Any]]:
    init_db()
    with get_connection() as conn:
        rows = conn.execute("SELECT * FROM review_term_bases ORDER BY updated_at DESC, filename").fetchall()
    return [_term_base_to_dict(row) for row in rows]


def get_term_base(term_base_id: str | None) -> dict[str, Any] | None:
    if not term_base_id:
        return None
    init_db()
    with get_connection() as conn:
        row = conn.execute("SELECT * FROM review_term_bases WHERE id = ?", (term_base_id,)).fetchone()
    return _term_base_to_dict(row) if row else None


def upload_term_base(filename: str, data: bytes) -> dict[str, Any]:
    init_db()
    ensure_directories()
    safe_name = Path(str(filename or "")).name.strip()
    suffix = Path(safe_name).suffix.lower()
    if not safe_name or suffix not in SUPPORTED_TERM_BASE_EXTENSIONS:
        raise TermBaseError("术语表仅支持 .xlsx 或 .xlsm 格式")
    if not data:
        raise TermBaseError("术语表文件为空")
    if len(data) > MAX_TERM_BASE_BYTES:
        raise TermBaseError("术语表不能超过 50 MB")

    parsed = _parse_term_table_workbook(data)
    file_hash = hashlib.sha256(data).hexdigest()
    now = utc_now()
    with get_connection() as conn:
        existing = conn.execute(
            "SELECT * FROM review_term_bases WHERE filename = ? COLLATE NOCASE", (safe_name,)
        ).fetchone()
        term_base_id = str(existing["id"]) if existing else uuid.uuid4().hex
        old_path = Path(str(existing["stored_path"])) if existing else None
        stored_path = TERM_BASES_DIR / f"{term_base_id}_{file_hash[:16]}{suffix}"
        stored_path.write_bytes(data)
        if existing:
            conn.execute(
                """
                UPDATE review_term_bases
                SET filename = ?, stored_path = ?, file_hash = ?, size_bytes = ?, sheet_name = ?,
                    entry_count = ?, languages_json = ?, has_entry_note = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    safe_name, str(stored_path), file_hash, len(data), parsed["sheet_name"],
                    len(parsed["entries"]), dumps_json(parsed["languages"]),
                    1 if parsed["has_entry_note"] else 0, now, term_base_id,
                ),
            )
            conn.execute("DELETE FROM review_term_entries WHERE term_base_id = ?", (term_base_id,))
        else:
            conn.execute(
                """
                INSERT INTO review_term_bases (
                    id, filename, stored_path, file_hash, size_bytes, sheet_name, entry_count,
                    languages_json, has_entry_note, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    term_base_id, safe_name, str(stored_path), file_hash, len(data), parsed["sheet_name"],
                    len(parsed["entries"]), dumps_json(parsed["languages"]),
                    1 if parsed["has_entry_note"] else 0, now, now,
                ),
            )
        conn.executemany(
            """
            INSERT INTO review_term_entries (
                id, term_base_id, row_number, entry_id, entry_subject, entry_note,
                terms_json, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    uuid.uuid4().hex, term_base_id, entry["row_number"], entry["entry_id"],
                    entry["entry_subject"], entry["entry_note"], dumps_json(entry["terms"]), now, now,
                )
                for entry in parsed["entries"]
            ],
        )

    _clear_matcher_cache(term_base_id)
    if old_path and old_path != stored_path:
        try:
            old_path.resolve().relative_to(TERM_BASES_DIR.resolve())
            old_path.unlink(missing_ok=True)
        except (OSError, ValueError):
            pass
    return get_term_base(term_base_id) or {}


def delete_term_base(term_base_id: str) -> None:
    term_base = get_term_base(term_base_id)
    if not term_base:
        raise TermBaseError("术语表不存在")
    with get_connection() as conn:
        conn.execute("UPDATE review_sessions SET term_base_id = NULL WHERE term_base_id = ?", (term_base_id,))
        conn.execute("DELETE FROM review_term_entries WHERE term_base_id = ?", (term_base_id,))
        conn.execute("DELETE FROM review_term_bases WHERE id = ?", (term_base_id,))
    _clear_matcher_cache(term_base_id)
    path = Path(str(term_base.get("stored_path") or ""))
    try:
        path.resolve().relative_to(TERM_BASES_DIR.resolve())
        path.unlink(missing_ok=True)
    except (OSError, ValueError):
        pass


def match_terms(
    term_base_id: str | None,
    source_language: str,
    target_language: str,
    source_text: str,
    *,
    limit: int = MAX_MATCHED_ENTRIES_PER_ITEM,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Return only term pairs whose source-language term occurs in source_text."""
    term_base = get_term_base(term_base_id)
    if not term_base or not str(source_text or "").strip():
        return [], {"status": "inactive"}
    available = list(term_base.get("languages") or [])
    source_column = resolve_term_table_language(source_language, available)
    target_column = resolve_term_table_language(target_language, available)
    if not source_column or not target_column:
        return [], {
            "status": "language_unmapped",
            "source_column": source_column,
            "target_column": target_column,
            "available_languages": available,
        }
    matcher = _get_matcher(term_base, source_column, target_column)
    normalized_text = _normalize_match_text(source_text)
    hits: dict[str, tuple[int, int, str]] = {}
    trie = matcher["trie"]
    for start in range(len(normalized_text)):
        node = trie
        cursor = start
        while cursor < len(normalized_text) and normalized_text[cursor] in node:
            node = node[normalized_text[cursor]]
            cursor += 1
            for pattern in node.get("_patterns", ()):  # longest variants win later
                if pattern["boundary"] and not _has_word_boundaries(normalized_text, start, cursor):
                    continue
                previous = hits.get(pattern["entry_key"])
                score = cursor - start
                if previous is None or score > previous[1] - previous[0]:
                    hits[pattern["entry_key"]] = (start, cursor, pattern["source"])
    ordered = sorted(hits.items(), key=lambda item: (item[1][0], -(item[1][1] - item[1][0])))[: max(1, limit)]
    matched = []
    for entry_key, (_, _, source_term) in ordered:
        entry = matcher["entries"][entry_key]
        pair: dict[str, Any] = {"source": source_term, "targets": entry["target_terms"]}
        if entry["entry_note"]:
            pair["entry_note"] = entry["entry_note"]
        matched.append(pair)
    return matched, {
        "status": "matched",
        "source_column": source_column,
        "target_column": target_column,
        "term_base_hash": term_base["file_hash"],
        "truncated": len(hits) > len(matched),
    }


def resolve_term_table_language(language: str, available: list[str]) -> str:
    normalized_available = {_normalize_language_key(item): item for item in available}
    key = _normalize_language_key(language)
    if key in normalized_available:
        return normalized_available[key]
    for candidate in LANGUAGE_ALIASES.get(key, ()):
        exact = normalized_available.get(_normalize_language_key(candidate))
        if exact:
            return exact
    return ""


def _parse_term_table_workbook(data: bytes) -> dict[str, Any]:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise TermBaseError(f"无法读取术语表：{exc}") from exc
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        raw_headers = next(rows, None)
        if not raw_headers:
            raise TermBaseError("术语表没有表头")
        headers = [_cell_text(value) for value in raw_headers]
        header_meta = [_metadata_header(value) for value in headers]
        if "entry_id" not in header_meta:
            raise TermBaseError("术语表缺少 Entry_ID 表头")
        language_indexes: dict[str, list[int]] = {}
        languages: list[str] = []
        for index, header in enumerate(headers):
            if not header or header_meta[index]:
                continue
            if header not in language_indexes:
                language_indexes[header] = []
                languages.append(header)
            language_indexes[header].append(index)
        if not languages:
            raise TermBaseError("术语表没有可识别的语种列")
        metadata_indexes = {
            name: header_meta.index(name) for name in METADATA_HEADERS.values() if name in header_meta
        }
        entries: list[dict[str, Any]] = []
        for row_number, values in enumerate(rows, 2):
            row = list(values)
            terms: dict[str, list[str]] = {}
            for language, indexes in language_indexes.items():
                variants = _dedupe(_cell_text(row[index]) if index < len(row) else "" for index in indexes)
                if variants:
                    terms[language] = variants
            if not terms:
                continue
            entries.append(
                {
                    "row_number": row_number,
                    "entry_id": _row_value(row, metadata_indexes.get("entry_id")),
                    "entry_subject": _row_value(row, metadata_indexes.get("entry_subject")),
                    "entry_note": _row_value(row, metadata_indexes.get("entry_note")),
                    "terms": terms,
                }
            )
        if not entries:
            raise TermBaseError("术语表没有可用术语条目")
        return {
            "sheet_name": sheet.title,
            "languages": languages,
            "has_entry_note": "entry_note" in metadata_indexes,
            "entries": entries,
        }
    finally:
        workbook.close()


def _get_matcher(term_base: dict[str, Any], source_column: str, target_column: str) -> dict[str, Any]:
    cache_key = (term_base["id"], term_base["file_hash"], source_column, target_column)
    with _matcher_lock:
        cached = _matcher_cache.get(cache_key)
        if cached is not None:
            _matcher_cache.move_to_end(cache_key)
            return cached
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT id, entry_note, terms_json FROM review_term_entries WHERE term_base_id = ? ORDER BY row_number",
            (term_base["id"],),
        ).fetchall()
    trie: dict[str, Any] = {}
    entries: dict[str, Any] = {}
    for row in rows:
        terms = loads_json(row["terms_json"], {})
        source_terms = _dedupe(terms.get(source_column) or [])
        target_terms = _dedupe(terms.get(target_column) or [])
        if not source_terms or not target_terms:
            continue
        entry_key = str(row["id"])
        entries[entry_key] = {"target_terms": target_terms, "entry_note": str(row["entry_note"] or "")}
        for source in source_terms:
            normalized = _normalize_match_text(source)
            if not normalized:
                continue
            node = trie
            for char in normalized:
                node = node.setdefault(char, {})
            node.setdefault("_patterns", []).append(
                {
                    "entry_key": entry_key,
                    "source": source,
                    "boundary": _needs_word_boundaries(normalized),
                }
            )
    matcher = {"trie": trie, "entries": entries}
    with _matcher_lock:
        _matcher_cache[cache_key] = matcher
        _matcher_cache.move_to_end(cache_key)
        while len(_matcher_cache) > 16:
            _matcher_cache.popitem(last=False)
    return matcher


def _clear_matcher_cache(term_base_id: str) -> None:
    with _matcher_lock:
        for key in [key for key in _matcher_cache if key[0] == term_base_id]:
            _matcher_cache.pop(key, None)


def _term_base_to_dict(row: Any) -> dict[str, Any]:
    return {
        "id": row["id"], "filename": row["filename"], "stored_path": row["stored_path"],
        "file_hash": row["file_hash"], "size_bytes": row["size_bytes"],
        "sheet_name": row["sheet_name"], "entry_count": row["entry_count"],
        "languages": loads_json(row["languages_json"], []),
        "has_entry_note": bool(row["has_entry_note"]),
        "created_at": row["created_at"], "updated_at": row["updated_at"],
    }


def _metadata_header(value: str) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().casefold()).strip("_")
    return METADATA_HEADERS.get(normalized, "")


def _normalize_language_key(value: str) -> str:
    return re.sub(r"[^0-9a-z\u3400-\u9fff]+", "", unicodedata.normalize("NFKC", str(value or "")).casefold())


def _normalize_match_text(value: str) -> str:
    return unicodedata.normalize("NFKC", str(value or "")).casefold()


def _needs_word_boundaries(value: str) -> bool:
    return bool(value and value[0].isascii() and value[-1].isascii() and value[0].isalnum() and value[-1].isalnum())


def _has_word_boundaries(text: str, start: int, end: int) -> bool:
    return (start == 0 or not text[start - 1].isalnum()) and (end == len(text) or not text[end].isalnum())


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _row_value(row: list[Any], index: int | None) -> str:
    return _cell_text(row[index]) if index is not None and index < len(row) else ""


def _dedupe(values: Any) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = str(value or "").strip()
        key = _normalize_match_text(text)
        if text and key not in seen:
            seen.add(key)
            result.append(text)
    return result
