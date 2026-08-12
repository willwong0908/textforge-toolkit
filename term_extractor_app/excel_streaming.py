"""Low-memory helpers for reading modern Excel workbooks.

The helpers in this module deliberately keep the workbook in ``read_only``
mode.  Callers must consume rows as an iterator instead of building a full
worksheet-sized list or DataFrame.
"""

from __future__ import annotations

from contextlib import contextmanager
from pathlib import Path
from typing import Iterator, Sequence

from openpyxl import load_workbook


STREAMABLE_EXTENSIONS = {".xlsx", ".xlsm"}


def is_streamable_excel(path: str | Path) -> bool:
    return Path(path).suffix.lower() in STREAMABLE_EXTENSIONS


@contextmanager
def open_streaming_workbook(path: str | Path, *, data_only: bool = True):
    """Open an xlsx/xlsm workbook without expanding all cells into memory."""
    workbook = load_workbook(path, read_only=True, data_only=data_only, keep_vba=Path(path).suffix.lower() == ".xlsm")
    try:
        for worksheet in workbook.worksheets:
            # Some exported files have a stale ``dimension`` value.  Reset it
            # before iteration so read-only mode discovers actual populated
            # cells instead of silently truncating columns.
            try:
                worksheet.reset_dimensions()
            except (AttributeError, ValueError):
                pass
        yield workbook
    finally:
        workbook.close()


def cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def header_map_from_values(values: Sequence[object]) -> dict[str, int]:
    """Map non-empty headers to their leftmost zero-based column index."""
    headers: dict[str, int] = {}
    for index, value in enumerate(values):
        name = cell_text(value).strip()
        if name and name not in headers:
            headers[name] = index
    return headers


def iter_sheet_rows(worksheet, *, min_row: int = 1, max_col: int | None = None, values_only: bool = True) -> Iterator[tuple[int, tuple[object, ...]]]:
    """Yield sheet rows with their one-based row numbers."""
    options = {"min_row": min_row, "values_only": values_only}
    if max_col is not None:
        options["max_col"] = max_col
    for row_number, row in enumerate(worksheet.iter_rows(**options), start=min_row):
        yield row_number, tuple(row)
