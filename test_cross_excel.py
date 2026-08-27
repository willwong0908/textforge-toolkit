from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from term_extractor_app.cross_excel import merge_excel_files_by_headers


class CrossExcelMergeTests(unittest.TestCase):
    def test_merge_keeps_styles_when_selected_columns_contain_empty_cells(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source_path = root / "source.xlsx"
            output_dir = root / "output"
            output_dir.mkdir()
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            sheet.append(["id", "fr", "kr"])
            sheet.append(["record-1", None, "안녕하세요"])
            styled_cell = sheet.cell(row=2, column=1)
            styled_cell.font = Font(bold=True, color="FFFFFF")
            styled_cell.fill = PatternFill("solid", fgColor="1F4E78")
            workbook.save(source_path)
            workbook.close()

            with patch("term_extractor_app.cross_excel._cross_excel_output_dir", return_value=output_dir):
                result = merge_excel_files_by_headers(
                    str(root),
                    ["fr", "id", "kr"],
                    apply_format=True,
                )

            self.assertEqual(result["row_count"], 1)
            merged = load_workbook(result["output_file"], data_only=False)
            try:
                merged_sheet = merged["合并数据"]
                self.assertIsNone(merged_sheet.cell(row=2, column=3).value)
                self.assertEqual(merged_sheet.cell(row=2, column=4).value, "record-1")
                self.assertEqual(merged_sheet.cell(row=2, column=5).value, "안녕하세요")
                self.assertTrue(merged_sheet.cell(row=2, column=4).font.bold)
                self.assertEqual(merged_sheet.cell(row=2, column=4).fill.fgColor.rgb, "001F4E78")
            finally:
                merged.close()


if __name__ == "__main__":
    unittest.main()
